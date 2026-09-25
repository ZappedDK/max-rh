"""
Gera Excel com lista de cargos e descrições para validação do RH.
"""
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

CARGOS = [
    ("AÇOUGUEIRO", "Responsável pelo corte, preparo e exposição de carnes bovinas, suínas e aves. Realiza o atendimento no balcão do açougue, orienta clientes sobre cortes e mantém a higiene e organização do setor."),
    ("ANALISTA COMERCIAL", "Analisa indicadores de vendas, acompanha o desempenho comercial das lojas e apoia estratégias de precificação e promoções. Elabora relatórios e apresenta resultados à gestão."),
    ("ANALISTA DE CADASTRO", "Responsável pelo cadastro e manutenção de fornecedores, produtos e clientes nos sistemas da empresa. Verifica dados, corrige inconsistências e garante a integridade das informações."),
    ("ANALISTA DE COMPRAS", "Realiza cotações, negocia condições com fornecedores e efetua a compra de mercadorias para as lojas. Acompanha pedidos, prazos de entrega e disponibilidade de produtos."),
    ("ANALISTA DE CREDITO", "Avalia solicitações de crédito de clientes e parceiros, analisa histórico financeiro e define limites. Monitora inadimplência e apoia a gestão de risco de crédito."),
    ("ANALISTA DE DEPARTAMENTO PESSOAL", "Executa rotinas de admissão, demissão, folha de pagamento, controle de ponto e benefícios. Garante o cumprimento da legislação trabalhista e suporte aos colaboradores."),
    ("ANALISTA DE ESTOQUE", "Monitora entradas e saídas de mercadorias, realiza inventários, identifica divergências e propõe melhorias nos processos de armazenagem e controle de estoque."),
    ("ANALISTA DE RECURSOS HUMANOS", "Apoia processos de recrutamento e seleção, treinamento, avaliação de desempenho e clima organizacional. Auxilia na gestão de talentos e desenvolvimento de pessoas."),
    ("ANALISTA FINANCEIRO", "Acompanha o fluxo de caixa, contas a pagar e receber, conciliações bancárias e relatórios financeiros. Apoia a tomada de decisões com base em indicadores e análises."),
    ("ANALISTA FISCAL", "Verifica e apura obrigações fiscais e tributárias, emite e valida notas fiscais, e garante o cumprimento das legislações tributárias aplicáveis ao negócio."),
    ("ASSISTENTE DE ESTOQUE", "Auxilia no controle de entrada e saída de mercadorias, organiza o estoque, separa produtos para as lojas e apoia na realização de inventários."),
    ("ASSISTENTE DE MONITORAMENTO", "Opera sistemas de câmeras e monitoramento das lojas. Registra ocorrências, alerta sobre situações suspeitas e apoia a equipe de prevenção de perdas."),
    ("ASSISTENTE FINANCEIRO", "Auxilia nas rotinas financeiras como lançamentos, conciliações, controle de pagamentos e recebimentos. Organiza documentos e apoia o analista financeiro."),
    ("ASSISTENTE FISCAL", "Apoia na conferência de notas fiscais, escrituração de documentos e geração de arquivos fiscais. Auxilia no cumprimento das obrigações acessórias."),
    ("ASSISTENTE JURIDICO", "Apoia a área jurídica na elaboração de documentos, acompanhamento de processos, pesquisa de legislação e organização de contratos e arquivos legais."),
    ("ASSISTENTE DE PREVENCAO DE PERDAS", "Auxilia na identificação e prevenção de perdas operacionais e de mercadorias. Apoia os processos de inventário e monitora indicadores de quebra."),
    ("ATENDENTE DE PADARIA", "Atende clientes no setor de padaria, organiza o espaço, embala produtos, repõe itens no balcão e garante a qualidade e apresentação dos produtos oferecidos."),
    ("AUXILIAR ADMINISTRATIVO", "Realiza atividades de apoio administrativo como digitação, arquivo de documentos, atendimento interno e organização de processos do setor."),
    ("AUXILIAR DE COMPRAS", "Apoia o processo de compras com cotações, organização de pedidos, controle de prazos e comunicação com fornecedores."),
    ("AUXILIAR DE CONFEITEIRO", "Auxilia na produção de bolos, doces, tortas e demais produtos de confeitaria, seguindo receitas e padrões de qualidade estabelecidos."),
    ("AUXILIAR DE DEPOSITO", "Realiza o recebimento, conferência, organização e movimentação de mercadorias no depósito da loja. Apoia nas atividades de carga e descarga."),
    ("AUXILIAR DE PADARIA", "Auxilia o padeiro na produção de pães e produtos de panificação, prepara insumos, limpa e organiza o ambiente de trabalho."),
    ("AUXILIAR FINANCEIRO", "Apoia as rotinas financeiras com lançamentos, controle de documentos e organização de processos. Auxilia na gestão de contas e relatórios básicos."),
    ("BALCONISTA DE AÇOUGUE", "Atende clientes no balcão do açougue, orienta sobre os cortes disponíveis, pesa e embala produtos, mantém a organização e higiene do setor."),
    ("CAPTADOR DE CLIENTE", "Aborda clientes nas proximidades das lojas, divulga promoções e produtos, incentiva o ingresso à loja e apoia ações de marketing e fidelização."),
    ("CONFEITEIRO", "Produz bolos, doces, tortas e sobremesas com criatividade e técnica. Garante a qualidade, apresentação e sabor dos produtos de confeitaria."),
    ("CONFERENTE", "Confere a entrada e saída de mercadorias, verifica quantidades e condições dos produtos recebidos e emite relatórios de divergências."),
    ("DESOSSADOR", "Especializado em separar a carne dos ossos de maneira eficiente, garantindo o aproveitamento máximo das peças e a qualidade dos cortes."),
    ("ENCARREGADO DE AÇOUGUE", "Gerencia as atividades do setor de açougue, coordena a equipe, controla estoque, garante a qualidade dos produtos e o atendimento ao cliente."),
    ("ENCARREGADO DE CHECK-OUT", "Supervisiona os caixas da loja, apoia operadores, resolve situações de atendimento, controla filas e garante a eficiência no check-out."),
    ("ENCARREGADO DE DEPOSITO", "Coordena as atividades do depósito, organiza o recebimento e armazenagem de mercadorias, gerencia a equipe e controla o fluxo de estoque."),
    ("ENCARREGADO DE FRIOS", "Responsável pela gestão do setor de frios, controla temperatura, validade e apresentação dos produtos, coordena a equipe e garante a qualidade."),
    ("ENCARREGADO DE HORTIFRUTI", "Gerencia o setor de frutas, verduras e legumes, controla qualidade, organiza exposição, coordena a equipe e garante o abastecimento adequado."),
    ("ENCARREGADO DE LOJA", "Apoia a gestão geral da loja, supervisiona setores, resolve problemas operacionais do dia a dia e garante o padrão de atendimento e organização."),
    ("ENCARREGADO DE PADARIA", "Coordena as atividades da padaria, escala a equipe, controla a produção, garante a qualidade dos produtos e o abastecimento do setor."),
    ("ENCARREGADO DE TESOURARIA", "Gerencia as atividades de tesouraria da loja, supervisiona o fechamento de caixa, controla sangrias, depósitos e a movimentação financeira."),
    ("ENCARREGADO FINANCEIRO", "Supervisiona as rotinas financeiras de uma unidade ou setor, controla fluxo de caixa, apoia conciliações e garante o cumprimento de metas."),
    ("ENCARREGADO DE PREVENCAO DE PERDAS", "Lidera a equipe de prevenção de perdas, coordena processos de inventário, monitora indicadores e implementa ações para redução de quebras e furtos."),
    ("ESTAGIARIO - ADMINISTRATIVO", "Apoia setores administrativos da empresa como RH, financeiro, compras ou marketing. Realiza atividades práticas vinculadas ao curso de graduação."),
    ("ESTAGIÁRIO - CAIXA", "Aprende e apoia as rotinas do setor de caixa, como atendimento ao cliente, operação de equipamentos e procedimentos de pagamento."),
    ("ESTAGIÁRIO - LOJA", "Apoia setores operacionais da loja como repositor, atendimento e organização, integrando teoria acadêmica à prática do varejo supermercadista."),
    ("FISCAL DE CAIXA", "Supervisiona os caixas em operação, autoriza descontos e cancelamentos, apoia operadores, garante a integridade dos valores e resolve ocorrências."),
    ("GERENTE DE LOJA", "Responsável pela gestão completa de uma unidade: pessoas, resultados, atendimento, abastecimento, prevenção de perdas e cumprimento de metas."),
    ("GERENTE DE MANUTENÇÃO", "Coordena a equipe de manutenção, gerencia contratos de serviços, planeja manutenções preventivas e corretivas e controla o patrimônio das lojas."),
    ("MOTORISTA", "Realiza transporte de mercadorias, colaboradores ou documentos com segurança. Cuida da conservação do veículo e cumpre prazos e rotas definidos."),
    ("MOTORISTA CARRETEIRO", "Conduz carretas para transporte de grandes volumes de mercadorias entre centros de distribuição e lojas, cumprindo normas de trânsito e segurança."),
    ("OPERADOR DE CAIXA", "Realiza o atendimento no caixa, processa pagamentos em dinheiro, cartão e outros meios, efetua sangrias e mantém a organização do posto de trabalho."),
    ("OPERADOR DE CARTAO", "Apoia nas operações com cartão de crédito e débito, administra máquinas POS, confere transações e oferece suporte ao cliente em dúvidas sobre pagamentos."),
    ("OPERADOR DE EMPILHADEIRA", "Opera empilhadeiras para movimentação de paletes e mercadorias no depósito, garantindo segurança, organização e agilidade nas operações."),
    ("PADEIRO", "Produz pães, roscas e produtos de panificação, controla fermentação e forno, garante sabor, textura e apresentação adequados aos padrões da loja."),
    ("PRECIFICADOR", "Realiza a etiquetagem e precificação de produtos nas gôndolas, confere preços no sistema, corrige divergências e mantém a loja organizada e sinalizada."),
    ("RECEPCIONISTA", "Recebe e direciona visitantes, atende telefone, organiza agendamentos e presta informações gerais. Representa a empresa no primeiro contato presencial."),
    ("REPOSITOR", "Abastece as gôndolas com mercadorias, organiza e rotaciona produtos, verifica validades e mantém a loja apresentável e com estoque adequado."),
    ("REPOSITOR DE FRIOS", "Responsável pelo abastecimento e organização do setor de frios, verifica validade e temperatura dos produtos, garantindo qualidade e apresentação."),
    ("REPOSITOR DE HORTIFRUTI", "Abastece e organiza o setor de frutas, verduras e legumes, seleciona produtos em bom estado, descarta avariados e mantém o setor sempre atrativo."),
    ("SUBGERENTE", "Apoia o gerente de loja na gestão das operações, lidera equipes na ausência do gerente, resolve ocorrências e garante o padrão de atendimento."),
    ("SUPERVISOR DE ESTOQUE", "Supervisiona o controle e a gestão do estoque das lojas, coordena inventários, analisa divergências e propõe melhorias nos processos de armazenagem."),
]

# ── Estilos ──────────────────────────────────────────────
COR_AZUL_ESCURO = "1E3A5F"   # cabeçalho
COR_AZUL_MEDIO  = "2563EB"   # destaque
COR_BRANCO      = "FFFFFF"
COR_CINZA_CLARO = "F1F5F9"   # linhas pares
COR_CINZA_BORDA = "CBD5E1"
COR_AMARELO     = "FEF9C3"   # coluna validação
COR_VERDE       = "DCFCE7"   # aprovado
COR_VERMELHO    = "FEE2E2"   # reprovado

borda_thin = Border(
    left=Side(style='thin', color=COR_CINZA_BORDA),
    right=Side(style='thin', color=COR_CINZA_BORDA),
    top=Side(style='thin', color=COR_CINZA_BORDA),
    bottom=Side(style='thin', color=COR_CINZA_BORDA),
)

def header_style(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=True, color=COR_BRANCO, size=11, name="Calibri")
    cell.fill = PatternFill("solid", fgColor=COR_AZUL_ESCURO)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = borda_thin
    return cell

wb = openpyxl.Workbook()

# ══════════════════════════════════════════════════════════
# ABA 1 — Lista de Cargos e Descrições
# ══════════════════════════════════════════════════════════
ws = wb.active
ws.title = "Cargos e Descrições"

# Título geral
ws.merge_cells("A1:E1")
titulo = ws["A1"]
titulo.value = "MAX SUPERMERCADOS — VALIDAÇÃO DE CARGOS E DESCRIÇÕES"
titulo.font = Font(bold=True, color=COR_BRANCO, size=13, name="Calibri")
titulo.fill = PatternFill("solid", fgColor=COR_AZUL_MEDIO)
titulo.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

# Subtítulo
ws.merge_cells("A2:E2")
sub = ws["A2"]
sub.value = "Documento gerado automaticamente pelo Sistema de RH — Para uso interno"
sub.font = Font(italic=True, color="64748B", size=9, name="Calibri")
sub.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 18

ws.row_dimensions[3].height = 6  # espaçamento

# Cabeçalho da tabela
headers = ["Nº", "Cargo", "Descrição das Atividades", "Status RH", "Observações"]
for col, h in enumerate(headers, 1):
    header_style(ws, 4, col, h)
ws.row_dimensions[4].height = 26

# Larguras das colunas
ws.column_dimensions["A"].width = 5
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 68
ws.column_dimensions["D"].width = 18
ws.column_dimensions["E"].width = 30

# Linhas de dados
for i, (cargo, desc) in enumerate(CARGOS, start=1):
    row = i + 4
    bg = COR_CINZA_CLARO if i % 2 == 0 else COR_BRANCO

    # Nº
    c_num = ws.cell(row=row, column=1, value=i)
    c_num.font = Font(size=10, color="94A3B8", name="Calibri")
    c_num.alignment = Alignment(horizontal="center", vertical="center")
    c_num.fill = PatternFill("solid", fgColor=bg)
    c_num.border = borda_thin

    # Cargo
    c_cargo = ws.cell(row=row, column=2, value=cargo)
    c_cargo.font = Font(bold=True, size=10, name="Calibri", color="1E293B")
    c_cargo.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    c_cargo.fill = PatternFill("solid", fgColor=bg)
    c_cargo.border = borda_thin

    # Descrição
    c_desc = ws.cell(row=row, column=3, value=desc)
    c_desc.font = Font(size=10, name="Calibri", color="334155")
    c_desc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    c_desc.fill = PatternFill("solid", fgColor=bg)
    c_desc.border = borda_thin

    # Status RH (campo para preenchimento — dropdown-like hint)
    c_status = ws.cell(row=row, column=4, value="Pendente")
    c_status.font = Font(size=10, name="Calibri", color="92400E")
    c_status.alignment = Alignment(horizontal="center", vertical="center")
    c_status.fill = PatternFill("solid", fgColor=COR_AMARELO)
    c_status.border = borda_thin

    # Observações (vazio para preenchimento)
    c_obs = ws.cell(row=row, column=5, value="")
    c_obs.fill = PatternFill("solid", fgColor=bg)
    c_obs.border = borda_thin
    c_obs.alignment = Alignment(vertical="center", wrap_text=True, indent=1)

    ws.row_dimensions[row].height = 42

# Rodapé
last_row = len(CARGOS) + 5 + 1
ws.merge_cells(f"A{last_row}:E{last_row}")
rod = ws[f"A{last_row}"]
rod.value = f"Total: {len(CARGOS)} cargos  |  Legenda — Status: Aprovado / Reprovado / Ajustar / Pendente"
rod.font = Font(italic=True, size=9, color="64748B", name="Calibri")
rod.alignment = Alignment(horizontal="right", vertical="center")
ws.row_dimensions[last_row].height = 18

# Congelar painel no cabeçalho
ws.freeze_panes = "A5"

# ══════════════════════════════════════════════════════════
# ABA 2 — Legenda de Status
# ══════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Legenda")
ws2.column_dimensions["A"].width = 20
ws2.column_dimensions["B"].width = 50

ws2.merge_cells("A1:B1")
tit2 = ws2["A1"]
tit2.value = "LEGENDA — STATUS DE VALIDAÇÃO"
tit2.font = Font(bold=True, color=COR_BRANCO, size=12, name="Calibri")
tit2.fill = PatternFill("solid", fgColor=COR_AZUL_ESCURO)
tit2.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 28

legenda = [
    ("Aprovado",  "Cargo e descrição validados pelo RH — pode ir ao ar",           "DCFCE7", "166534"),
    ("Reprovado", "Cargo ou descrição não correspondem à realidade — não incluir",  "FEE2E2", "991B1B"),
    ("Ajustar",   "Descrição necessita de correção ou complemento pelo RH",         "FEF9C3", "92400E"),
    ("Pendente",  "Aguardando análise do RH (padrão ao gerar o documento)",         "F1F5F9", "475569"),
]

header_style(ws2, 2, 1, "Status")
header_style(ws2, 2, 2, "Significado")
ws2.row_dimensions[2].height = 22

for i, (status, sig, bg, fg) in enumerate(legenda, start=3):
    c1 = ws2.cell(row=i, column=1, value=status)
    c1.font = Font(bold=True, size=10, color=fg, name="Calibri")
    c1.fill = PatternFill("solid", fgColor=bg)
    c1.alignment = Alignment(horizontal="center", vertical="center")
    c1.border = borda_thin

    c2 = ws2.cell(row=i, column=2, value=sig)
    c2.font = Font(size=10, color="334155", name="Calibri")
    c2.fill = PatternFill("solid", fgColor=bg)
    c2.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    c2.border = borda_thin
    ws2.row_dimensions[i].height = 26

# ── Salvar ───────────────────────────────────────────────
output = r"c:\Users\wanderson.ferreira\Documents\programa python any\Gerenciador de Colaborador\MAX_Cargos_Validacao_RH.xlsx"
wb.save(output)
print(f"[OK] Arquivo gerado: {output}")
print(f"[OK] Total de cargos: {len(CARGOS)}")
